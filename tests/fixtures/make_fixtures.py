"""Генератор маленьких фикстур всех форматов (результат закоммичен в tests/fixtures/).

    uv run python tests/fixtures/make_fixtures.py

Текст — фрагменты «Исповеди» Л. Н. Толстого (общественное достояние).
"""

import zipfile
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).parent
P1 = "Я был крещён и воспитан в православной христианской вере. Меня учили ей и с детства, и во всё время моего отрочества и юности."
P2 = "Но когда я 18-ти лет вышел со второго курса университета, я не верил уже ничему из того, чему меня учили."
P3 = "Судя по некоторым воспоминаниям, я никогда и не верил серьёзно, а имел только доверие к тому, чему меня учили."
P4 = "Жизнь моя остановилась. Я мог дышать, есть, пить, спать и не мог не дышать, не есть, не пить, не спать; но жизни не было."
P5 = "Со мною сделалось то, что я стал просыпаться по ночам и спрашивать себя: зачем я живу?"


def fb2() -> bytes:
    return f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE lolz [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
<description><title-info>
  <author><first-name>Лев</first-name><middle-name>Николаевич</middle-name><last-name>Толстой</last-name></author>
  <book-title>Исповедь</book-title><lang>ru</lang><date>1882</date>
</title-info></description>
<body>
  <title><p>Исповедь</p></title>
  <section><title><p>Часть первая</p></title>
    <section><title><p>Глава I</p></title>
      <epigraph><p>Эпиграф к главе.</p><text-author>Автор эпиграфа</text-author></epigraph>
      <p>{P1}<a l:href="#n1" type="note">1</a></p>
      <p>{P2}</p>
      <p>&xxe;</p>
    </section>
    <section><title><p>Глава II</p></title>
      <p>{P3}</p>
      <poem><stanza><v>Строка стиха первая</v><v>Строка стиха вторая</v></stanza></poem>
    </section>
  </section>
  <section><title><p>Часть вторая</p></title>
    <section><title><p>Глава III</p></title><p>{P4}</p><p>{P5}</p></section>
  </section>
</body>
<body name="notes">
  <section id="n1"><title><p>1</p></title><p>Примечание о крещении.</p></section>
</body>
<binary id="cover.jpg" content-type="image/jpeg">AAAA</binary>
</FictionBook>
""".encode()


EPUB_CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""
EPUB_OPF = """<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
  <dc:identifier id="id">tolstoy-test</dc:identifier><dc:title>Исповедь</dc:title>
  <dc:creator>Лев Толстой</dc:creator><dc:language>ru</dc:language>
</metadata>
<manifest>
  <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
  <item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/>
  <item id="c1" href="text/ch1.xhtml" media-type="application/xhtml+xml"/>
  <item id="c2" href="text/ch2.xhtml" media-type="application/xhtml+xml"/>
</manifest>
<spine><itemref idref="cover" linear="no"/><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""
EPUB_NAV = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body>
<nav epub:type="toc"><ol>
  <li><a href="text/ch1.xhtml">Часть первая</a>
    <ol><li><a href="text/ch1.xhtml#s1">Глава I</a></li></ol></li>
  <li><a href="text/ch2.xhtml">Глава II</a></li>
</ol></nav></body></html>"""


def xhtml(title: str, *paras: str) -> str:
    body = "".join(f"<p>{p}</p>" for p in paras)
    return (
        '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
        f"<body><h1>{title}</h1>{body}<ul><li>Пункт списка</li></ul></body></html>"
    )


def epub(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", EPUB_CONTAINER)
        z.writestr("OEBPS/content.opf", EPUB_OPF)
        z.writestr("OEBPS/nav.xhtml", EPUB_NAV)
        z.writestr("OEBPS/cover.xhtml", xhtml("Обложка", "Текст обложки"))
        z.writestr("OEBPS/text/ch1.xhtml", xhtml("Глава I", P1, P2))
        z.writestr("OEBPS/text/ch2.xhtml", xhtml("Глава II", P3, P4))


def pdf(path: Path) -> None:
    """6 страниц: колонтитул и номер на каждой, перенос слова, outline из двух глав.

    Текст — через insert_htmlbox: у base-14 шрифтов PyMuPDF нет кириллицы в извлечении.
    """
    doc = pymupdf.open()
    pages = [
        ("Глава первая", [P1, "Это слово с перено-<br>сом внутри абзаца."]),
        (None, [P2]),
        (None, [P3]),
        ("Глава вторая", [P4]),
        (None, [P5]),
        (None, ["Последняя страница книги."]),
    ]
    for i, (heading, paras) in enumerate(pages, start=1):
        page = doc.new_page()
        page.insert_htmlbox(pymupdf.Rect(72, 30, 520, 60), "<p>Л. Н. Толстой. Исповедь</p>")
        y = 100
        if heading:
            page.insert_htmlbox(pymupdf.Rect(72, y, 520, y + 30), f"<h2>{heading}</h2>")
            y += 50
        for para in paras:
            page.insert_htmlbox(pymupdf.Rect(72, y, 520, y + 140), f"<p>{para}</p>")
            y += 160
        page.insert_htmlbox(pymupdf.Rect(290, 790, 320, 810), f"<p>{i}</p>")
    doc.set_toc([[1, "Глава первая", 1], [1, "Глава вторая", 4]])
    doc.set_metadata({"title": "Исповедь", "author": "Лев Толстой"})
    doc.save(path)


def scan_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for _ in range(3):
        page = doc.new_page()
        page.draw_rect(pymupdf.Rect(50, 50, 500, 700), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    doc.save(path)


def docx_file(path: Path) -> None:
    d = docx.Document()
    d.core_properties.title = "Исповедь"
    d.core_properties.author = "Лев Толстой"
    d.add_paragraph("Исповедь", style="Title")
    d.add_heading("Глава I", level=1)
    d.add_paragraph(P1)
    d.add_paragraph(P2)
    d.add_heading("Раздел 1.1", level=2)
    d.add_paragraph(P3)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Год"
    table.cell(0, 1).text = "Событие"
    table.cell(1, 0).text = "1879"
    table.cell(1, 1).text = "Начало работы над «Исповедью»"
    d.add_heading("Глава II", level=1)
    d.add_paragraph(P4)
    d.save(path)


def txt() -> str:
    return f"ИСПОВЕДЬ\n\nГЛАВА I\n\n{P1}\n{P2}\n\n{P3}\n\n* * *\n\nГЛАВА II\n\n{P4}\n\n{P5}\n"


def main() -> None:
    (HERE / "tolstoy.fb2").write_bytes(fb2())
    with zipfile.ZipFile(HERE / "tolstoy.fb2.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("tolstoy.fb2", fb2())
    epub(HERE / "tolstoy.epub")
    pdf(HERE / "tolstoy.pdf")
    scan_pdf(HERE / "scan.pdf")
    docx_file(HERE / "tolstoy.docx")
    (HERE / "chapters_cp1251.txt").write_bytes(txt().encode("cp1251"))
    (HERE / "broken.pdf").write_bytes(b"%PDF-1.7\n" + b"\x00garbage" * 50)


if __name__ == "__main__":
    main()
